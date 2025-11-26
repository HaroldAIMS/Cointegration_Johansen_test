#!/usr/bin/env python3
"""
Script de conversion des rapports Markdown en documents Word (.docx)
Utilise python-docx pour générer des documents formatés
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re


def parse_markdown_to_docx(md_file, docx_file):
    """
    Convertit un fichier Markdown en document Word formaté
    """
    doc = Document()
    
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    lines = content.split('\n')
    current_table = None
    
    for i, line in enumerate(lines):
        line_stripped = line.strip()
        
        # Sauter les lignes vides
        if not line_stripped:
            current_table = None
            continue
        
        # Titre H1 (# Title)
        if line_stripped.startswith('# ') and not line_stripped.startswith('## '):
            text = line_stripped.lstrip('# ').strip()
            p = doc.add_heading(text, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            current_table = None
        
        # Titre H2 (## Subtitle)
        elif line_stripped.startswith('## ') and not line_stripped.startswith('### '):
            text = line_stripped.lstrip('# ').strip()
            p = doc.add_heading(text, level=2)
            current_table = None
        
        # Titre H3 (### Subtitle)
        elif line_stripped.startswith('### '):
            text = line_stripped.lstrip('# ').strip()
            p = doc.add_heading(text, level=3)
            current_table = None
        
        # Listes (- ou *)
        elif line_stripped.startswith(('- ', '* ', '+ ')):
            text = line_stripped.lstrip('-* +').strip()
            doc.add_paragraph(text, style='List Bullet')
            current_table = None
        
        # Listes numérotées (1. 2. etc)
        elif re.match(r'^\d+\.\s', line_stripped):
            text = re.sub(r'^\d+\.\s', '', line_stripped)
            doc.add_paragraph(text, style='List Number')
            current_table = None
        
        # Tables Markdown (| col1 | col2 |)
        elif line_stripped.startswith('|'):
            # Détecter si c'est une ligne de séparateur
            if all(c in '|-: ' for c in line_stripped):
                continue
            
            # Parser les colonnes
            cols = [c.strip() for c in line_stripped.split('|')[1:-1]]
            
            if cols:
                # Créer nouvelle table ou ajouter ligne
                if current_table is None:
                    current_table = doc.add_table(rows=1, cols=len(cols))
                    current_table.style = 'Light Grid Accent 1'
                    
                    # Header row
                    for j, col in enumerate(cols):
                        current_table.rows[0].cells[j].text = col
                else:
                    # Ajouter ligne de données
                    row = current_table.add_row()
                    for j, col in enumerate(cols):
                        row.cells[j].text = col
        
        # Séparateur horizontal ---
        elif line_stripped == '---' or line_stripped == '***':
            doc.add_paragraph('_' * 60)
            current_table = None
        
        # Paragraphe normal
        else:
            if line_stripped:
                doc.add_paragraph(line_stripped)
            current_table = None
    
    # Sauvegarder
    doc.save(docx_file)
    print(f"✅ Converti : {Path(md_file).name} → {Path(docx_file).name}")


def main():
    """
    Convertit tous les rapports Markdown en documents Word
    """
    notebooks_dir = Path('/Users/harold/DataAnalyctisandScience/Cointegration_Johansen_test/Inflation Notebooks')
    
    # Fichiers Markdown à convertir
    md_files = sorted(notebooks_dir.glob('*_REPORT.md'))
    
    print(f"🔄 Conversion de {len(md_files)} rapports Markdown en docx...\n")
    
    success = 0
    for md_file in md_files:
        docx_file = md_file.with_suffix('.docx')
        try:
            parse_markdown_to_docx(str(md_file), str(docx_file))
            success += 1
        except Exception as e:
            print(f"❌ Erreur lors de la conversion de {md_file.name}: {e}")
    
    print(f"\n✨ Conversion terminée ! {success}/{len(md_files)} fichiers .docx mis à jour")


if __name__ == '__main__':
    main()
