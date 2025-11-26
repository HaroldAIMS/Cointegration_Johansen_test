#!/usr/bin/env python3
"""
Script de conversion Markdown → Word (.docx)
Utilise python-docx pour générer des documents professionnels
"""

import sys
import os
from pathlib import Path

# Vérifier et installer si nécessaire
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Installation de python-docx...")
    os.system(f"{sys.executable} -m pip install -q python-docx")
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

import re


def parse_markdown(md_file):
    """Parse fichier Markdown et retourne liste de blocs formatés"""
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    blocks = []
    lines = content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Skip lignes vides
        if not line.strip():
            i += 1
            continue
        
        # Déterminer le type de bloc
        stripped = line.strip()
        
        # Titres
        if stripped.startswith('### '):
            blocks.append(('h3', stripped.lstrip('# ').strip()))
        elif stripped.startswith('## '):
            blocks.append(('h2', stripped.lstrip('# ').strip()))
        elif stripped.startswith('# '):
            blocks.append(('h1', stripped.lstrip('# ').strip()))
        
        # Tables
        elif stripped.startswith('|'):
            if all(c in '|-: ' for c in stripped):
                # Skip séparateurs
                pass
            else:
                cols = [c.strip() for c in stripped.split('|')[1:-1]]
                blocks.append(('table_row', cols))
        
        # Listes
        elif stripped.startswith(('- ', '* ', '+ ')):
            text = stripped.lstrip('-*+ ').strip()
            blocks.append(('bullet', text))
        
        # Listes numérotées
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            blocks.append(('number', text))
        
        # Séparateur horizontal
        elif stripped == '---' or stripped == '***' or stripped == '___':
            blocks.append(('hr', None))
        
        # Paragraphes normaux
        elif stripped:
            blocks.append(('paragraph', stripped))
        
        i += 1
    
    return blocks


def write_docx(blocks, output_file):
    """Écrit les blocs dans un document Word"""
    doc = Document()
    
    # Style par défaut
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    
    table_in_progress = False
    current_table = None
    
    for block_type, content in blocks:
        
        if block_type == 'h1':
            p = doc.add_heading(content, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Couleur bleu foncé
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 51, 102)
        
        elif block_type == 'h2':
            doc.add_heading(content, level=2)
        
        elif block_type == 'h3':
            doc.add_heading(content, level=3)
        
        elif block_type == 'table_row':
            if not table_in_progress:
                # Nouvelle table
                current_table = doc.add_table(rows=1, cols=len(content))
                current_table.style = 'Light Grid Accent 1'
                table_in_progress = True
                
                # Header
                for i, col in enumerate(content):
                    cell = current_table.rows[0].cells[i]
                    cell.text = col
            else:
                # Row supplémentaire
                row = current_table.add_row()
                for i, col in enumerate(content):
                    row.cells[i].text = col
        
        elif block_type == 'bullet':
            doc.add_paragraph(content, style='List Bullet')
        
        elif block_type == 'number':
            doc.add_paragraph(content, style='List Number')
        
        elif block_type == 'hr':
            doc.add_paragraph('_' * 50)
        
        elif block_type == 'paragraph':
            p = doc.add_paragraph(content)
            # Appliquer formatage simple
            for run in p.runs:
                run.font.size = Pt(11)
        
        # Reset table si on quitte le mode table
        if block_type != 'table_row':
            table_in_progress = False
    
    doc.save(output_file)


def main():
    notebooks_dir = Path('/Users/harold/DataAnalyctisandScience/Cointegration_Johansen_test/Inflation Notebooks')
    
    # Trouver tous les fichiers Markdown à convertir
    md_files = sorted([f for f in notebooks_dir.glob('*_REPORT.md')])
    
    print(f"\n📝 Conversion de {len(md_files)} rapports Markdown en docx")
    print("=" * 50)
    
    success_count = 0
    
    for md_file in md_files:
        try:
            # Parser et écrire
            blocks = parse_markdown(str(md_file))
            docx_file = md_file.with_suffix('.docx')
            write_docx(blocks, str(docx_file))
            
            print(f"✅ {md_file.name:30} ({len(blocks)} blocs)")
            success_count += 1
            
        except Exception as e:
            print(f"❌ {md_file.name:30} Erreur: {str(e)[:30]}")
    
    print("=" * 50)
    print(f"\n✨ Résultat: {success_count}/{len(md_files)} fichiers convertis")
    print(f"📂 Localisation: {notebooks_dir}\n")


if __name__ == '__main__':
    main()
